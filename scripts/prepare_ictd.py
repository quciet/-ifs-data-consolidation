"""Source-specific GRD 2025 preprocessing. Originals/base are read-only.

Requires openpyxl for extraction, plus this repository's pinned DataGator lookup.
The reviewed profile is deliberately restricted to the supplied 2025 workbook.
No merging, interpolation, rounding, or territorial arithmetic is performed.
"""
import argparse
import collections
import csv
import hashlib
import json
import math
from pathlib import Path
import shutil
import sqlite3
import sys
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

def q(s):
    return '"' + s.replace('"', '""') + '"'

def digest(p):
    h=hashlib.sha256()
    with Path(p).open('rb') as f:
        for chunk in iter(lambda:f.read(4*1024*1024),b''):h.update(chunk)
    return h.hexdigest()

def write_json(p,v):
    Path(p).write_text(json.dumps(v,indent=2,ensure_ascii=False,default=str)+'\n',encoding='utf-8')

def write_csv(p,rows,fields):
    with Path(p).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)

def year_value(v):
    if isinstance(v,bool):raise ValueError('Boolean year')
    s=str(v)
    if len(s)!=4 or not s.isascii() or not s.isdigit() or not 1980<=int(s)<=2024:
        raise ValueError(f'Unexpected GRD 2025 year {v!r}')
    return int(s)

def observation(v,location,approved_errors):
    if v is None:return None,'source_null'
    if isinstance(v,str) and not v.strip():return None,'empty_string' if v=='' else 'whitespace'
    if location in approved_errors and v=='#VALUE!':return None,'approved_excel_error'
    if isinstance(v,bool) or not isinstance(v,(int,float)) or not math.isfinite(v):
        raise ValueError(f'Unsupported observation {location}: {v!r}')
    return float(v)*100.0,'fraction_to_percent'

def make_specs(old):
    columns={'Grants':('BD','Grants'),'Property':('AN','Property taxes'),
             'Resources':('U','Total resource revenue'),'SocSec':('BC','Social contributions'),
             'NonTax':('AZ','Total non-tax revenue'),'TaxDirect':('AB','Direct taxes including social contributions and resource taxes'),
             'TaxGoodSer':('AR','Total taxes on goods and services'),'TaxICorporate':('AJ','Total corporate income taxes'),
             'TaxIncome':('AF','Total taxes on income, profits and capital gains'),'TaxIndirect':('AO','Total indirect taxes'),
             'TaxIndividual':('AI','Personal income taxes'),'TaxTot':('W','Total taxes including social contributions'),
             'TaxTrade':('AV','Total taxes on international trade')}
    result=[]
    for r in old:
        t=r['Table'];sheet='General' if 'Gen' in t or 'Calc' in t else 'Central'
        if 'Calc' in t or 'CurRev' in t:col,label='Q','Total revenue including grants and social contributions'
        else:
            stem=t.split('Gen' if sheet=='General' else 'Cen',1)[1].split('%')[0]
            if stem=='NonTaxTot':stem='NonTax'
            col,label=columns[stem]
        result.append({'table':t,'variable':r['Variable'],'sheet':sheet,'column':col,'label':label})
    return result

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--repo',type=Path,required=True);ap.add_argument('--source',type=Path,required=True);ap.add_argument('--base',type=Path,required=True);ap.add_argument('--output',type=Path,help='Explicit output override; default is a new versioned folder inside --source');ap.add_argument('--profile',type=Path,required=True)
    a=ap.parse_args();a.repo=a.repo.resolve();a.source=a.source.resolve();a.base=a.base.resolve()
    profile=json.loads(a.profile.read_text(encoding='utf-8'))
    for name,h in profile['source_sha256'].items():assert digest(a.source/name)==h, f'Source changed: {name}'
    sys.path.insert(0,str(a.repo/'src'))
    from ifs_pipeline.concordance import lookup_index,normalize_name
    from ifs_pipeline.preparation_paths import create_preparation_workspace
    a.output,input_files=create_preparation_workspace(a.source,a.output)
    evidence=a.output/'Working Files';evidence.mkdir();raw=evidence/'Original download';raw.mkdir()
    input_hashes={str(f):digest(f) for f in input_files}
    for f,h in input_hashes.items():
        dest=raw/Path(f).relative_to(a.source);dest.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(f,dest);assert digest(dest)==h
    base_hashes={str(a.base/n):digest(a.base/n) for n in ['DataDict.db','IFsHistSeries.db']}
    shutil.copy2(a.profile,evidence/'ictd.profile.json');shutil.copy2(__file__,evidence/'prepare_ictd.py')
    shutil.copytree(a.repo/'reference/datagator',evidence/'datagator')
    hist=sqlite3.connect((a.base/'IFsHistSeries.db').as_uri()+'?mode=ro',uri=True);hist.row_factory=sqlite3.Row
    bc=sqlite3.connect((a.base/'DataDict.db').as_uri()+'?mode=ro',uri=True);bc.row_factory=sqlite3.Row
    canonical={r['FIPS_CODE']:r['Country'] for r in hist.execute('select Country,FIPS_CODE from SeriesPopulation')};assert len(canonical)==188
    index,_=lookup_index(json.loads((evidence/'datagator/country_data.json').read_text()),canonical)
    write_csv(evidence/'master.csv',[{'Country':n,'FIPS_CODE':c} for c,n in canonical.items()],['Country','FIPS_CODE'])
    frame=list(load_workbook(raw/'IFsFrame.xlsx',read_only=True,data_only=True).active.values)
    assert {(r[0],r[1]) for r in frame[1:]}=={(n,c) for c,n in canonical.items()}
    ow=load_workbook(raw/'DataDict ICTD_Old 20260921.xlsx',read_only=True,data_only=True);rr=list(ow.active.values);old=[dict(zip(rr[0],r)) for r in rr[1:] if any(v is not None for v in r)];ow.close()
    specs=make_specs(old);assert len(specs)==27
    write_json(evidence/'table_mapping.json',specs)
    schema=list(bc.execute('pragma table_info(DataDict)'));fields=[r['name'] for r in schema]
    schema_sql='CREATE TABLE DataDict ('+', '.join(q(r['name'])+' '+r['type'] for r in schema)+')'
    (evidence/'datadict_schema.sql').write_text(schema_sql+';\n',encoding='utf-8')
    base_meta={r['Table']:dict(r) for r in bc.execute('select * from DataDict') if r['Table'] in {s['table'] for s in specs}}
    assert len(base_meta)==27
    write_json(evidence/'base_metadata.json',base_meta);write_json(evidence/'old_workbook_metadata.json',old)
    w=load_workbook(raw/'UNUWIDERGRD_2025.xlsx',read_only=True,data_only=True)
    source_rows={};identities=[];exclusions=[];annotations=[];key_notes=[]
    for sn in ['General','Central']:
        rows=list(w[sn].values);seen=set();source_rows[sn]={}
        write_json(evidence/(sn+'_headers.json'),rows[:3])
        for n,r in enumerate(rows[3:],4):
            if all(v is None for v in r):continue
            y=year_value(r[7]);name=r[2];iso=r[6];ident=r[0]
            choices=index.get(normalize_name(name),set());code=next(iter(choices)) if len(choices)==1 else None
            method='pinned_datagator_alias'
            if name in profile['name_overrides']:
                code=profile['name_overrides'][name];method='reviewed_source_name_and_iso'
            if sn=='General' and n==1515:
                assert (ident,name,iso,y)==('ETH1980','Estonia','ETH',1980)
                assert r[3]=='Sub-Saharan Africa'
                code='ETH';method='reviewed_row_identity_correction'
            if code is None:
                assert name in profile['excluded_names'],f'Unresolved identity {sn}:{n} {name}'
                exclusions.append({'sheet':sn,'excel_row':n,'source_country':name,'source_iso':iso,'year':y,'raw_row_json':json.dumps(r,ensure_ascii=False),'reason':'Outside selected base canonical roster; no geographic aggregation'})
                continue
            expected_iso=profile['ifs_to_iso'].get(code,code)
            assert iso==expected_iso,(sn,n,name,iso,code)
            assert str(ident).startswith(iso),(sn,n,ident,iso)
            if str(ident)!=iso+str(y):
                key_notes.append({'sheet':sn,'excel_row':n,'source_country':name,'identifier':ident,'year':y,'reason':'Use explicit year column; identifier is not the observation year authority','source_note':r[8]})
            k=(code,y);assert k not in seen, f'Duplicate country/year {sn}:{k}';seen.add(k)
            source_rows[sn][k]=(n,r)
            identities.append({'sheet':sn,'excel_row':n,'original_name':name,'source_iso':iso,'identifier':ident,'year':y,'Country':canonical[code],'FIPS_CODE':code,'method':method})
            if any(isinstance(v,str) and v.strip() or isinstance(v,(int,float)) and v!=0 for v in r[8:16]):
                annotations.append({'sheet':sn,'excel_row':n,'Country':canonical[code],'FIPS_CODE':code,'year':y,'source_annotation_json':json.dumps(r[8:16],ensure_ascii=False)})
    approved_errors={tuple(x) for x in profile['approved_error_cells']}
    db=a.output/'IFsDataImport_UNUWIDER_ICTD_2025.db';conn=sqlite3.connect(db);conn.execute(schema_sql)
    audit=[];missing=[];review=[];changes=[];coverage=[];screen=[];metadata_out={};expected={};actual_errors=[]
    for spec in specs:
        t=spec['table'];sn=spec['sheet'];ci=column_index_from_string(spec['column'])-1
        vals={};ys=set();countries=set()
        for (code,y),(n,r) in source_rows[sn].items():
            loc=(sn,canonical[code],y,spec['column']);v,kind=observation(r[ci],loc,approved_errors)
            record={'table':t,'sheet':sn,'cell':spec['column']+str(n),'original_country':r[2],'Country':canonical[code],'FIPS_CODE':code,'year':y,'raw_value_json':json.dumps(r[ci]),'operation':kind,'output_value':v}
            audit.append(record)
            if v is None:missing.append(record)
            else:
                vals[(code,y)]=v;ys.add(y);countries.add(code)
                if v<0 or v>100:screen.append({**record,'finding':'negative' if v<0 else 'above_100_percent','action':'retained_source_value'})
            if kind=='approved_excel_error':actual_errors.append(loc)
        assert ys
        years=list(range(min(y for c,y in source_rows[sn]),max(y for c,y in source_rows[sn])+1))
        cols=['Country','FIPS_CODE']+[str(y) for y in years]+['Earliest','MostRecent']
        conn.execute('CREATE TABLE '+q(t)+' ('+', '.join(q(c)+' '+('VARCHAR(255)' if i<2 else 'DOUBLE(53)') for i,c in enumerate(cols))+')')
        expected[t]={}
        for code,name in canonical.items():
            vv=[vals.get((code,y)) for y in years];nn=[v for v in vv if v is not None]
            row=[name,code]+vv+[nn[0] if nn else None,nn[-1] if nn else None]
            conn.execute('INSERT INTO '+q(t)+' VALUES ('+','.join('?' for _ in row)+')',row);expected[t][code]=row
        m=dict(base_meta[t]);level='General government' if sn=='General' else 'Central government'
        notes='Includes source caution flags and source-defined historical coverage; consult the GRD 2025 country notes. No territorial reallocation is applied.'
        if 'SocSec' in t:notes+=' Social-contribution coverage and classification vary across countries.'
        if 'Resources' in t:notes+=' Resource revenue is sparsely reported; missing values are not zero.'
        if 'CurRev' in t:notes+=' Legacy IFs identifier retained; the matched source measure is total central-government revenue including grants and social contributions.'
        if any(loc[0]==sn and loc[3]==spec['column'] for loc in approved_errors):notes+=' Source Excel errors are missing: '+('Madagascar 1993-2010.' if sn=='General' else 'Pakistan 2021-2023.')
        m.update({'Definition':level+' '+spec['label'].lower()+', as a percentage of GDP.',
                  'Extended Source Defn':'GRD 2025 expresses revenue using a common GDP series, predominantly IMF World Economic Outlook April 2025. Government level and inclusion of grants, social contributions and resource revenue follow the named source measure.',
                  'Years':f'{min(ys)}-{max(ys)}','Source':'UNU-WIDER Government Revenue Dataset',
                  'Original Source':'https://doi.org/10.35188/UNU-WIDER/GRD-2025','Notes':notes,
                  'Last IFs Update':profile['preparation_date'],'Name in Source':spec['label'],
                  'Country Concordance':'IFs Country','Formula':'Excel stored fraction * 100'})
        # All remaining fields retain the established, matching IFs 8.73 configuration.
        for field in fields:
            prev=base_meta[t][field];v=m[field]
            if field in ['CURRENCY','UsedInPreprocessorFileName','DisplayNotes'] and v is None:
                disposition='intentionally blank';why={'CURRENCY':'Percent of GDP; no output currency amount.','UsedInPreprocessorFileName':'Base usage flag is 0; no preprocessor module recorded.','DisplayNotes':'No distinct display annotation established; source caveats are in Notes.'}[field]
            elif v!=prev:disposition='supported value';why=profile['metadata_reasons'].get(field,'Source headers, GRD 2025 documentation and table-specific review.')
            else:disposition='preserve existing';why='Matching Table/Variable in supplied old workbook and selected IFs 8.73 DataDict; table-specific configuration retained.'
            review.append({'Table':t,'Variable':m['Variable'],'field':field,'proposed_value_json':json.dumps(v,ensure_ascii=False),'disposition':disposition,'evidence':why})
            if v!=prev:changes.append({'Table':t,'Variable':m['Variable'],'field':field,'base_json':json.dumps(prev,ensure_ascii=False),'prepared_json':json.dumps(v,ensure_ascii=False)})
        assert m['Variable']==spec['variable'] and m['Units']=='Percent'
        conn.execute('INSERT INTO DataDict ('+','.join(q(f) for f in fields)+') VALUES ('+','.join('?' for _ in fields)+')',[m[f] for f in fields]);metadata_out[t]=m
        coverage.append({'table':t,'sheet':sn,'source_column':spec['column'],'country_rows':len(canonical),'source_country_rows':len({c for c,y in source_rows[sn]}),'countries_with_observations':len(countries),'first_observed_year':min(ys),'last_observed_year':max(ys),'year_columns':len(years),'observations':len(vals),'zero_values':sum(v==0 for v in vals.values()),'missing_grid_cells':len(canonical)*len(years)-len(vals),'source_error_cells':sum(x[0]==sn and x[3]==spec['column'] for x in actual_errors),'status':'ready_with_findings'})
    assert set(actual_errors)==approved_errors and len(actual_errors)==len(approved_errors)
    conn.commit();assert conn.execute('pragma integrity_check').fetchone()[0]=='ok';conn.close();w.close()
    # Independent saved-file reread, including every NULL, endpoint, metadata field and declared type.
    with sqlite3.connect(db.as_uri()+'?mode=ro',uri=True) as check:
        check.row_factory=sqlite3.Row
        assert len(list(check.execute('select * from DataDict')))==27
        for spec in specs:
            t=spec['table'];rows=list(check.execute('select * from '+q(t)));assert len(rows)==188
            for row in rows:assert list(row)==expected[t][row['FIPS_CODE']]
            assert dict(check.execute('select * from DataDict where "Table"=?',(t,)).fetchone())==metadata_out[t]
            declared=list(check.execute('pragma table_info('+q(t)+')'));assert all(r['type']==('VARCHAR(255)' if i<2 else 'DOUBLE(53)') for i,r in enumerate(declared))
        assert [(r['name'],r['type']) for r in check.execute('pragma table_info(DataDict)')]==[(r['name'],r['type']) for r in schema]
    for name,rows in [('observation_audit',audit),('missingness',missing),('metadata_review',review),('metadata_changes',changes),('coverage',coverage),('screening_findings',screen),('country_identity_review',identities),('excluded_rows',exclusions),('source_annotations',annotations),('year_identifier_notes',key_notes)]:
        if rows:write_csv(evidence/(name+'.csv'),rows,list(rows[0]))
    write_json(evidence/'reviewed_metadata.json',metadata_out)
    write_json(evidence/'recipe.json',{'source':'UNU-WIDER GRD 2025','adapter':'prepare_ictd.py','input_sha256':digest(raw/'UNUWIDERGRD_2025.xlsx'),'mapping':specs,'multiplier':100,'rounding':False,'missing_values':['SQL NULL','empty string','whitespace','24 specifically approved #VALUE! cells'],'country_policy':profile,'base':str(a.base),'source_year_column':'H; integer or four-digit numeric text','readiness':'Preprocessing checked; separate validation remains next stage'})
    assert all(digest(Path(f))==h for f,h in input_hashes.items());assert all(digest(Path(f))==h for f,h in base_hashes.items())
    manifest={'status':'completed_with_issues','stage':'preprocessing','source_files':input_hashes,'base_files':base_hashes,'output':{'path':str(db),'sha256':digest(db)},'tables':coverage,'verified':{'saved_cells_metadata_schema':True,'source_and_base_unchanged':True,'integrity_check':'ok'},'excel_errors_to_null':len(actual_errors),'identity_corrections':1,'source_screening_flags':len(screen),'next_stage':'validation','merge_performed':False}
    write_json(evidence/'preparation_manifest.json',manifest)
    write_json(evidence/'handoff.json',{'status':'completed_with_issues','stage_id':'preprocessing','scope':'27 Table/Variable entries in the supplied old DataDict workbook','base':str(a.base),'inputs':[{'path':f,'sha256':h,'role':'source'} for f,h in input_hashes.items()],'outputs':[{'path':str(db),'sha256':digest(db),'role':'prepared_import'}],'next_stage':'validation','table_outcomes':[{'table':s['table'],'outcome':'ready_with_findings','artifact':str(db)} for s in specs],'native_results':{'preparation_manifest':str(evidence/'preparation_manifest.json')}})
    report=['# UNU-WIDER ICTD preprocessing','',f'Prepared {len(specs)} tables and {sum(r["observations"] for r in coverage):,} numeric observations against IFs 8.73 (20260814). Each table has 188 canonical country rows.','',f'Import: `{db.name}`. No merge was performed. The base and original downloads are unchanged.','','## Source and scope','The complete supplied UNUWIDERGRD_2025.xlsx is the numeric authority. Central and General sheets remain separate. The short workbook is an incomplete subset and is retained as context. The old DataDict supplies the 27-table scope; matching base settings supply IFs configuration.','Publisher citation: https://doi.org/10.35188/UNU-WIDER/GRD-2025. Source reference: https://www.wider.unu.edu/database/data-and-resources-grd. The source Info sheet contains a stale 2023 citation; workbook title and supplied November 2025 guide establish the 2025 release.','','## Transformations and verification','Excel fraction values are multiplied by 100 to store Percent (e.g., 0.30223417816262066 becomes 30.223417816262067). Stored observations are not rounded to display precision. Blank/missing cells stay NULL, zero remains a value. Earliest and MostRecent copy existing first/last non-null values.','Table mappings, source cell addresses and raw values are retained in Working Files. Saved database values, NULLs, schemas and all DataDict fields were reread and checked. Country/year duplicates fail preparation.','','## Findings',f'- {len(actual_errors)} #VALUE! cells were converted to SQL NULL under the user\'s explicit 2026-09-22 instruction: Madagascar general social contributions, 1993-2010; Pakistan central personal and corporate income taxes, 2021-2023. Every token is preserved in missingness.csv.', '- General row 1515 labels ETH1980 as Estonia; its ETH code, identifier, Sub-Saharan region and Ethiopia GDP context identify Ethiopia. This source-specific correction is audited.', '- Source identifier/year differences are retained in year_identifier_notes.csv; explicit year fields govern. Bhutan\'s shifted years have source explanations. Congo 2021-2022 retain the published years and values, with stale identifiers flagged.',f'- {len(exclusions):,} country-year rows outside the base roster were excluded with complete raw row evidence; no territorial allocation or aggregation was performed.',f'- {len(screen):,} negative or above-100% observations were retained and flagged; source caution flags and historical-country caveats remain available in source_annotations.csv.', '- Total/direct tax mappings include social contributions. The two revenue tables use total revenue including grants and social contributions: General for GovtCalcRevTot, Central for GovtCurRev. This is supported by supplied labels, the short extract and comparisons with established baseline values; the legacy current-revenue identifier is retained with a clarified definition.', '- General/Central source headers place the goods-and-services heading one cell early; column AR is the Total column, confirmed by hierarchy and component values. Column AQ belongs to non-resource indirect taxes.', '- This adapter uses the supplied files only and performs no online fetching. No replacement observations are obtained or invented.','','## Coverage','| Table | Countries with data | Years | Observations |','|---|---:|---|---:|']
    report += [f'| {r["table"]} | {r["countries_with_observations"]} | {r["first_observed_year"]}-{r["last_observed_year"]} | {r["observations"]:,} |' for r in coverage]
    report += ['','Next stage: validation of the prepared import and source findings. This package does not certify a merged delivery.','']
    (a.output/'Preparation Report.md').write_text('\n'.join(report),encoding='utf-8')
    print(json.dumps(manifest,indent=2))

if __name__=='__main__':main()
